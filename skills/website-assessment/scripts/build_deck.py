#!/usr/bin/env python3
"""
build_deck.py - Stage 3. Turns findings.json + framed screenshots into the
Crackwits "Website UX/UI Assessment" PowerPoint deck.

Slide types produced:
  cover    - client name, URL, date
  section  - eyebrow + title, findings panel (left), framed screenshot with
             numbered markers (right)
  summary  - narrative paragraphs + category count grid
  priority - optional: critical/moderate findings ranked (when severity is used)

Every element is a native PowerPoint shape, so the deck stays fully editable.

Usage:
  python3 build_deck.py --findings audit/findings.json --root audit --out assessment.pptx
"""

import argparse
import json
from datetime import date
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

ROOT = Path(__file__).parent.parent
BRAND = json.loads((ROOT / "assets" / "brand.json").read_text())

SLIDE_W_IN = 13.3333
SLIDE_H_IN = 7.5
EMU_PER_IN = 914400
SW = int(SLIDE_W_IN * EMU_PER_IN)
SH = int(SLIDE_H_IN * EMU_PER_IN)

C = BRAND["colors"]
L = BRAND["layout"]
T = BRAND["type"]
CATS = BRAND["categories"]
SEV = BRAND["severity"]

# Toggled from findings.json meta; lets a deck run denser by dropping
# the optional per-finding lines.
SHOW = {"fix": True, "principle": True}


def rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def fx(f):  # fraction of slide width -> EMU
    return int(SW * f)


def fy(f):
    return int(SH * f)


# ------------------------------------------------------------------ text math
# Urbanist is fairly narrow. These constants were tuned against the reference
# slides; adjust CHAR_W_RATIO if wrapping looks off after a real run.
CHAR_W_RATIO = 0.485      # average glyph advance / font size
LINE_H_RATIO = 1.42


def wrapped_lines(text, font_pt, box_w_in):
    """Estimate how many lines `text` needs inside a box of box_w_in inches."""
    if not text:
        return 0
    char_w_in = (font_pt * CHAR_W_RATIO) / 72.0
    per_line = max(8, int(box_w_in / char_w_in))
    lines = 0
    for para in str(text).split("\n"):
        words, cur = para.split(), 0
        if not words:
            lines += 1
            continue
        n = 1
        for w in words:
            add = len(w) + (1 if cur else 0)
            if cur + add > per_line:
                n += 1
                cur = len(w)
            else:
                cur += add
        lines += n
    return lines


def text_h_in(text, font_pt, box_w_in):
    return wrapped_lines(text, font_pt, box_w_in) * font_pt * LINE_H_RATIO / 72.0


# ------------------------------------------------------------------ primitives

def _strip_theme_style(shape):
    """python-pptx autoshapes carry a <p:style> ref to the theme, which makes
    PowerPoint/LibreOffice paint a default drop shadow. We set every fill and
    line explicitly, so drop the theme reference entirely."""
    el = shape._element
    for tag in ("style",):
        node = el.find("{http://schemas.openxmlformats.org/presentationml/2006/main}" + tag)
        if node is not None:
            el.remove(node)
    spPr = el.spPr
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    if spPr.find(ns + "effectLst") is None:
        from lxml import etree
        spPr.append(etree.SubElement(spPr, ns + "effectLst"))


def add_rect(slide, x, y, w, h, fill=None, radius=None, line=None, line_w=None, shadow=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, x, y, w, h)
    if radius:
        # adjustment is radius / min(w, h)
        try:
            s.adjustments[0] = min(0.5, radius / max(1, min(w, h)))
        except Exception:
            pass
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w or 0.75)
    else:
        s.line.fill.background()
    if not shadow:
        s.shadow.inherit = False
        _strip_theme_style(s)
    s.text_frame.word_wrap = True
    return s


def add_text(slide, x, y, w, h, runs, size, color, font=None, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0, wrap=True):
    """runs: str, or list of (text, {bold, color, size}) tuples, or list of paragraphs."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    _strip_theme_style(tb)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor

    paras = runs if isinstance(runs, list) and runs and isinstance(runs[0], list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if i:
            p.space_before = Pt(size * 0.55)
        items = para if isinstance(para, list) else [(para, {})]
        for txt, opt in items:
            r = p.add_run()
            r.text = str(txt)
            f = r.font
            f.name = opt.get("font", font or T["finding_body"]["font"])
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.color.rgb = rgb(opt.get("color", color))
    return tb


def _avg_char_ratio(text):
    """Uppercase runs are noticeably wider than the mixed-case average."""
    if not text:
        return CHAR_W_RATIO
    caps = sum(1 for c in text if c.isupper())
    frac = caps / len(text)
    return CHAR_W_RATIO * (1 + 0.30 * frac)


def pill_w_in(text, size_pt, pad_x_in=0.058):
    return (len(text) * size_pt * _avg_char_ratio(text) / 72.0) + pad_x_in * 2


def add_pill(slide, x, y, text, color, size_pt, pad_x_in=0.058, h_in=None):
    h_in = h_in or L.get("tag_h_in", 0.145)
    w = Emu(int(pill_w_in(text, size_pt, pad_x_in) * EMU_PER_IN))
    h = Emu(int(h_in * EMU_PER_IN))
    s = add_rect(slide, x, y, w, h, fill=color, radius=int(h * 0.30))
    tf = s.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = T["tag"]["font"]
    r.font.size = Pt(size_pt)
    r.font.bold = False
    r.font.color.rgb = rgb(C["tag_text"])
    return s, w


def add_marker(slide, cx, cy, n, d_emu=None):
    d = d_emu or fx(L["marker_d"])
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(cx - d / 2), int(cy - d / 2), d, d)
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(C["marker"])
    s.line.color.rgb = rgb("#FFFFFF")
    s.line.width = Pt(1.1)
    s.shadow.inherit = False
    _strip_theme_style(s)
    tf = s.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = f"{n:02d}"
    r.font.name = T["marker"]["font"]
    r.font.size = Pt(T["marker"]["size_pt"])
    r.font.color.rgb = rgb(C["marker_text"])
    return s


# ------------------------------------------------------------------ chrome
def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = add_rect(s, 0, 0, SW, SH, fill=C["background"])
    bg.line.fill.background()
    return s


def add_footer(slide, meta):
    yr = meta.get("year") or date.today().year
    add_text(slide, fx(L["footer_l_x"]), fy(L["footer_y"]), fx(0.40), fy(0.03),
             BRAND["agency"]["footer_left"].format(year=yr),
             T["footer"]["size_pt"], C["footer"], font=T["footer"]["font"])
    add_text(slide, fx(L["footer_r_x"] - 0.30), fy(L["footer_y"]), fx(0.30), fy(0.03),
             meta.get("footer_right", BRAND["agency"]["footer_right"]),
             T["footer"]["size_pt"], C["footer"], font=T["footer"]["font"],
             align=PP_ALIGN.RIGHT)


def add_heading(slide, eyebrow, title):
    add_text(slide, fx(L["eyebrow"]["x"]), fy(L["eyebrow"]["y"]), fx(0.30), fy(0.045),
             eyebrow, T["eyebrow"]["size_pt"], C["eyebrow"], font=T["eyebrow"]["font"])
    add_text(slide, fx(L["title"]["x"]), fy(L["title"]["y"]), fx(0.30), fy(0.075),
             title, T["title"]["size_pt"], C["title"], font=T["title"]["font"])


# ------------------------------------------------------------------ slides
def cover_slide(prs, meta):
    s = blank(prs)
    add_text(s, fx(0.058), fy(0.36), fx(0.70), fy(0.06),
             meta.get("deck_kicker", "Website UX/UI Assessment"),
             T["eyebrow"]["size_pt"], C["eyebrow"], font=T["eyebrow"]["font"])
    add_text(s, fx(0.058), fy(0.42), fx(0.80), fy(0.16),
             meta.get("client", "Client"), 54, C["title"], font=T["title"]["font"])
    sub = meta.get("url", "")
    when = meta.get("audited_on") or date.today().isoformat()
    add_text(s, fx(0.058), fy(0.60), fx(0.70), fy(0.05),
             f"{sub}    ·    {when}", 12, C["footer"], font=T["footer"]["font"])
    add_footer(s, meta)
    return s


def tag_rows(f, avail_w_in, show_severity=True):
    """Lay tags (+ severity) into rows that fit avail_w_in. Returns list of rows."""
    items = [(c, CATS.get(c, {}).get("color", C["body_muted"]), T["tag"]["size_pt"])
             for c in f.get("categories", [])]
    sev = f.get("severity")
    if show_severity and sev in SEV:
        items.append((sev.upper(), SEV[sev]["color"], T["tag"]["size_pt"] - 0.5))
    gap = 0.030
    rows, cur, used = [], [], 0.0
    for label, col, sz in items:
        w = pill_w_in(label, sz)
        if cur and used + gap + w > avail_w_in:
            rows.append(cur)
            cur, used = [], 0.0
        cur.append((label, col, sz, w))
        used += w + (gap if len(cur) > 1 else 0)
    if cur:
        rows.append(cur)
    return rows


def _card_height_in(f, body_w_in, show_severity=True):
    """Height a finding card needs, in inches."""
    body = T["finding_body"]["size_pt"]
    tag_h = L.get("tag_h_in", 0.145)
    nrows = max(1, len(tag_rows(f, body_w_in, show_severity)))
    h = nrows * tag_h + (nrows - 1) * 0.028 + 0.055      # tag rows + gap to body
    h += text_h_in(f.get("observation", ""), body, body_w_in)
    if f.get("detail"):
        h += 0.07 + text_h_in(f["detail"], body, body_w_in)
    if f.get("fix") and SHOW.get("fix", True):
        h += 0.06 + text_h_in("Fix: " + f["fix"], body, body_w_in)
    if f.get("principle") and SHOW.get("principle", True):
        h += 0.04 + text_h_in("Principle: " + f["principle"], body, body_w_in)
    if f.get("scope"):
        h += 0.07 + text_h_in(f["scope"], body, body_w_in)
    return h + (L["card_pad_y"] * SLIDE_H_IN * 2)


def section_slide(prs, meta, page, section, img_path, findings, start_n, cont=False,
                  screen_rect=None):
    meta = {**meta, "_screen_rect": screen_rect}
    s = blank(prs)
    add_heading(s, page, section + (" (cont.)" if cont else ""))
    add_footer(s, meta)

    # ---- right: framed screenshot + markers
    dev = L["device"]
    if img_path and Path(img_path).exists():
        with Image.open(img_path) as im:
            iw, ih = im.size
        box_w, box_h = fx(dev["w"]), fy(dev["h"])
        scale = min(box_w / iw, box_h / ih)
        w, h = int(iw * scale), int(ih * scale)
        x = fx(dev["x"]) + (box_w - w) // 2
        y = fy(dev["y"]) + (box_h - h) // 2
        s.shapes.add_picture(str(img_path), x, y, w, h)
        img_rect = (x, y, w, h)
    else:
        img_rect = (fx(dev["x"]), fy(dev["y"]), fx(dev["w"]), fy(dev["h"]))
        ph = add_rect(s, *img_rect, fill="#3E3C3E", radius=fx(L["device_radius"]))
        add_text(s, img_rect[0], img_rect[1] + img_rect[3] // 2, img_rect[2], fy(0.04),
                 "screenshot unavailable", 11, C["footer"], align=PP_ALIGN.CENTER)

    # ---- left: findings panel
    p = L["panel"]
    px, py, pw, ph_ = fx(p["x"]), fy(p["y"]), fx(p["w"]), fy(p["h"])
    add_rect(s, px, py, pw, ph_, fill=C["panel"], radius=fx(L["panel_radius"]))

    pad = fx(L["panel_pad"])
    inner_x, inner_w = px + pad, pw - pad * 2
    cy = py + pad
    body_pt = T["finding_body"]["size_pt"]
    card_pad_x, card_pad_y = fx(L["card_pad_x"]), fy(L["card_pad_y"])
    text_x = inner_x + card_pad_x + fx(L["marker_d"]) + fx(0.008)
    text_w = inner_w - card_pad_x * 2 - fx(L["marker_d"]) - fx(0.008)
    text_w_in = text_w / EMU_PER_IN

    for i, f in enumerate(findings):
        n = start_n + i
        ch = int(_card_height_in(f, text_w_in, meta.get('show_severity', True)) * EMU_PER_IN)
        add_rect(s, inner_x, cy, inner_w, ch, fill=C["card"], radius=fx(L["card_radius"]))

        # marker chip mirrored in the card
        add_marker(s, inner_x + card_pad_x + fx(L["marker_d"]) / 2,
                   cy + card_pad_y + fx(L["marker_d"]) / 2, n)

        # category tags (wrapped into rows that fit the card)
        ty = cy + card_pad_y
        tag_h = L.get("tag_h_in", 0.145)
        rows = tag_rows(f, text_w_in, meta.get("show_severity", True))
        for r_i, row in enumerate(rows):
            tx = text_x
            ry = ty + int(r_i * (tag_h + 0.028) * EMU_PER_IN)
            for label, col, sz, w in row:
                add_pill(s, tx, ry, label, col, sz)
                tx += int((w + 0.030) * EMU_PER_IN)

        # body
        by = ty + int((len(rows) * tag_h + (len(rows) - 1) * 0.028 + 0.055) * EMU_PER_IN)
        paras = [[(f.get("observation", ""), {})]]
        if f.get("detail"):
            paras.append([(f["detail"], {})])
        if f.get("fix") and SHOW.get("fix", True):
            paras.append([("Fix: ", {"bold": True}), (f["fix"], {})])
        if f.get("principle") and SHOW.get("principle", True):
            paras.append([("Principle: ", {"bold": True}),
                          (f["principle"], {"color": C["body_muted"]})])
        if f.get("scope"):
            paras.append([(f["scope"], {"bold": True})])
        add_text(s, text_x, by, text_w, ch - (by - cy) - card_pad_y,
                 paras, body_pt, C["body"], font=T["finding_body"]["font"], spacing=1.16)

        # marker on the screenshot.
        # marker x/y are fractions of the LIVE SCREEN, not of the framed PNG,
        # so map them through screen_rect (bezel + shadow padding) when present.
        m = f.get("marker")
        if m and img_rect:
            sr = meta.get("_screen_rect") or {"x": 0.0, "y": 0.0, "w": 1.0, "h": 1.0}
            fxr = sr["x"] + sr["w"] * float(m.get("x", 0.5))
            fyr = sr["y"] + sr["h"] * float(m.get("y", 0.5))
            mx = img_rect[0] + img_rect[2] * fxr
            my = img_rect[1] + img_rect[3] * fyr
            add_marker(s, mx, my, n)

        cy += ch + fy(L["card_gap"])

    return s


def paginate(findings, panel_h_in, text_w_in, per_slide_cap=7, show_sev=True):
    """Split findings so cards always fit the panel."""
    pages, cur, used = [], [], 0.0
    for f in findings:
        h = _card_height_in(f, text_w_in, show_sev) + L["card_gap"] * SLIDE_H_IN
        if cur and (used + h > panel_h_in - 0.06 or len(cur) >= per_slide_cap):
            pages.append(cur)
            cur, used = [], 0.0
        cur.append(f)
        used += h
    if cur:
        pages.append(cur)
    return pages


def summary_slide(prs, meta, counts, narrative):
    s = blank(prs)
    add_heading(s, "General", "Summary")
    add_footer(s, meta)

    dev = L["device"]
    hero = meta.get("summary_image")
    if hero and Path(hero).exists():
        with Image.open(hero) as im:
            iw, ih = im.size
        box_w, box_h = fx(dev["w"] + 0.06), fy(dev["h"])
        scale = min(box_w / iw, box_h / ih)
        s.shapes.add_picture(str(hero), fx(dev["x"]), fy(dev["y"]),
                             int(iw * scale), int(ih * scale))

    p = L["panel"]
    px, py, pw, ph_ = fx(p["x"]), fy(p["y"]), fx(p["w"]), fy(p["h"])
    add_rect(s, px, py, pw, ph_, fill=C["panel"], radius=fx(L["panel_radius"]))

    pad = fx(L["panel_pad"])
    inner_x, inner_w = px + pad + fx(0.006), pw - (pad + fx(0.006)) * 2
    body_pt = T["finding_body"]["size_pt"]
    inner_w_in = inner_w / EMU_PER_IN

    ty = py + pad + fy(0.012)
    paras = [[(t, {})] for t in narrative]
    th = sum(text_h_in(t, body_pt, inner_w_in) for t in narrative) + 0.20 * len(narrative)
    add_text(s, inner_x, ty, inner_w, int(th * EMU_PER_IN), paras,
             body_pt, C["body"], font=T["finding_body"]["font"], spacing=1.20)

    # ---- 2-column count grid, sized so every used category fits
    ordered = [c for c in sorted(CATS, key=lambda k: CATS[k]["summary_order"]) if counts.get(c)]
    cols, gap = 2, fx(0.010)
    rows_n = max(1, -(-len(ordered) // cols))
    cw = (inner_w - gap) // cols

    grid_top = ty + int(th * EMU_PER_IN) + fy(0.022)
    grid_bottom = py + ph_ - pad
    avail = grid_bottom - grid_top
    chh = min(fy(0.078), max(fy(0.048), (avail - gap * (rows_n - 1)) // rows_n))
    # bottom-align the grid the way the reference slide does
    grid_h = rows_n * chh + gap * (rows_n - 1)
    gy = max(grid_top, grid_bottom - grid_h)

    for i, cat in enumerate(ordered):
        r, cidx = divmod(i, cols)
        x = inner_x + cidx * (cw + gap)
        y = gy + r * (chh + gap)
        add_rect(s, x, y, cw, chh, fill=C["card"], radius=fx(L["card_radius"]))
        add_pill(s, x + fx(0.010), y + (chh - int(L.get("tag_h_in", 0.145) * EMU_PER_IN)) // 2,
                 cat, CATS[cat]["color"], T["tag"]["size_pt"])
        add_text(s, x + cw - fx(0.058), y, fx(0.046), chh,
                 str(counts[cat]), T["stat_number"]["size_pt"], C["body"],
                 font=T["stat_number"]["font"], align=PP_ALIGN.RIGHT,
                 anchor=MSO_ANCHOR.MIDDLE)
    return s


def priority_slide(prs, meta, ranked):
    s = blank(prs)
    add_heading(s, "General", "Priority Fixes")
    add_footer(s, meta)

    p = L["panel"]
    px, py = fx(0.056), fy(0.244)
    pw, ph_ = fx(0.888), fy(0.670)
    add_rect(s, px, py, pw, ph_, fill=C["panel"], radius=fx(L["panel_radius"]))

    pad = fx(L["panel_pad"])
    x, y = px + pad, py + pad
    inner_w = pw - pad * 2
    body_pt = T["finding_body"]["size_pt"]

    for i, f in enumerate(ranked, 1):
        rowh = fy(0.058)
        if y + rowh > py + ph_ - pad:
            break
        add_rect(s, x, y, inner_w, rowh, fill=C["card"], radius=fx(L["card_radius"]))
        add_text(s, x + fx(0.012), y, fx(0.030), rowh, f"{i:02d}",
                 body_pt + 2, C["body_muted"], font=T["stat_number"]["font"],
                 anchor=MSO_ANCHOR.MIDDLE)
        sev = f.get("severity", "Moderate")
        add_pill(s, x + fx(0.044), y + int(rowh * 0.30), sev.upper(),
                 SEV.get(sev, SEV["Moderate"])["color"], T["tag"]["size_pt"] - 0.5)
        cat = (f.get("categories") or ["UX Design"])[0]
        add_pill(s, x + fx(0.100), y + int(rowh * 0.30), cat,
                 CATS.get(cat, {}).get("color", C["body_muted"]), T["tag"]["size_pt"])
        loc = f.get("_where", "")
        add_text(s, x + fx(0.210), y, fx(0.60), rowh,
                 [[(f.get("fix") or f.get("observation", ""), {}),
                   ((f"   — {loc}" if loc else ""), {"color": C["body_muted"]})]],
                 body_pt, C["body"], font=T["finding_body"]["font"],
                 anchor=MSO_ANCHOR.MIDDLE)
        y += rowh + fy(0.008)
    return s


# ------------------------------------------------------------------ main
def build(findings_path, root, out_path):
    data = json.loads(Path(findings_path).read_text())
    meta = data.get("meta", {})
    SHOW["fix"] = meta.get("show_fix", True)
    SHOW["principle"] = meta.get("show_principle", True)
    root = Path(root)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SW), Emu(SH)

    if meta.get("cover", True):
        cover_slide(prs, meta)

    counts, ranked = {}, []
    p = L["panel"]
    panel_h_in = p["h"] * SLIDE_H_IN
    text_w_in = (fx(p["w"]) - fx(L["panel_pad"]) * 2 - fx(L["card_pad_x"]) * 2
                 - fx(L["marker_d"]) - fx(0.008)) / EMU_PER_IN

    for sl in data.get("slides", []):
        img = sl.get("framed") or sl.get("screenshot")
        img_path = (root / img) if img and not Path(img).is_absolute() else (Path(img) if img else None)
        fs = sl.get("findings", [])
        for f in fs:
            for c in f.get("categories", []):
                counts[c] = counts.get(c, 0) + 1
            if f.get("severity") in ("Critical", "Moderate"):
                ranked.append({**f, "_where": f"{sl.get('page','')} · {sl.get('section','')}"})

        pages = paginate(fs, panel_h_in, text_w_in, show_sev=meta.get('show_severity', True))
        n = 1
        for i, chunk in enumerate(pages):
            section_slide(prs, meta, sl.get("page", ""), sl.get("section", ""),
                          img_path, chunk, n, cont=(i > 0),
                          screen_rect=sl.get("screen_rect"))
            n += len(chunk)

    summ = data.get("summary", {})
    if summ:
        meta = {**meta, "summary_image": (root / summ["image"]) if summ.get("image") else meta.get("summary_image")}
        summary_slide(prs, meta, summ.get("counts") or counts, summ.get("narrative", []))

    if ranked and meta.get("priority_slide", True):
        ranked.sort(key=lambda f: -SEV.get(f.get("severity", "Minor"), SEV["Minor"])["weight"])
        priority_slide(prs, meta, ranked[:9])

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return {"out": str(out_path), "slides": len(prs.slides.__iter__.__self__._sldIdLst),
            "counts": counts, "total_findings": sum(counts.values())}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--findings", required=True)
    ap.add_argument("--root", default=".")
    ap.add_argument("--out", default="assessment.pptx")
    a = ap.parse_args()
    print(json.dumps(build(a.findings, a.root, a.out), indent=2))


if __name__ == "__main__":
    main()
